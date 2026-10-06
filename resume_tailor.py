from __future__ import annotations

import re
import platform
import shutil
import subprocess
import tempfile
import time
from collections import Counter
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import deepcopy
from difflib import SequenceMatcher
from html import escape as html_escape
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from pydantic import BaseModel, Field

from advanced_engine import SKILL_ALIASES, clean_text, extract_required_skills, extract_skills


SECTION_ALIASES = {
    "professional summary": "Professional Summary",
    "profile summary": "Professional Summary",
    "summary": "Professional Summary",
    "technical skills": "Technical Skills",
    "core skills": "Technical Skills",
    "skills": "Technical Skills",
    "professional experience": "Professional Experience",
    "work experience": "Professional Experience",
    "experience": "Professional Experience",
    "selected projects": "Selected Projects",
    "projects": "Selected Projects",
    "project experience": "Selected Projects",
    "education": "Education",
    "academic qualification": "Education",
    "certifications & professional learning": "Certifications",
    "certifications and professional learning": "Certifications",
    "certifications": "Certifications",
    "achievements": "Achievements",
}

FOCUS_STOPWORDS = {
    "about", "ability", "added", "advantage", "after", "also", "analysis", "analyst",
    "candidate", "clear", "company", "data", "description", "essential", "experience",
    "from", "have", "ideal", "including", "into", "job", "limited", "more", "other",
    "preferred", "required", "requirements", "role", "similar", "skills", "strong",
    "support", "team", "that", "their", "this", "through", "using", "with", "work",
    "years", "your",
}

TAILORABLE_SECTIONS = {
    "Profile",
    "Professional Summary",
    "Technical Skills",
    "Professional Experience",
    "Selected Projects",
}


class ResumeRewrite(BaseModel):
    line_id: str
    revised_text: str
    reason: str


class ResumeTailoringResponse(BaseModel):
    rewrites: list[ResumeRewrite] = Field(default_factory=list)
    strong_matches: list[str] = Field(default_factory=list)
    remaining_gaps: list[str] = Field(default_factory=list)
    questions: list[str] = Field(default_factory=list)
    change_summary: list[str] = Field(default_factory=list)


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_.")
    return cleaned or "candidate"


def _normalize_candidate_name(candidate_name: str) -> str:
    return re.sub(r"[_-]+", " ", candidate_name).strip()


def _is_contact_line(line: str) -> bool:
    return bool(re.search(
        r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|linkedin\.com|(?:\+?\d[\d ()-]{8,}\d)",
        line,
        re.IGNORECASE,
    ))


def _section_heading(line: str) -> str | None:
    normalized = clean_text(line).strip(" :-")
    for key, label in SECTION_ALIASES.items():
        if normalized == key:
            return label
    if (
        2 <= len(line.split()) <= 5
        and len(line) <= 55
        and line.upper() == line
        and not any(character.isdigit() for character in line)
    ):
        return line.title()
    return None


def _is_date_line(line: str) -> bool:
    return bool(re.search(
        r"\b(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
        r"jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|"
        r"dec(?:ember)?|present|(?:19|20)\d{2})\b",
        line,
        re.IGNORECASE,
    ))


def _is_labeled_line(line: str) -> bool:
    prefix, separator, _ = line.partition(":")
    return bool(separator and 1 <= len(prefix.split()) <= 5 and len(prefix) <= 38)


def _prepare_source_lines(source_resume: str) -> list[str]:
    """Join extraction-created line wraps without dropping any source wording."""
    prepared: list[str] = []
    for raw_line in source_resume.splitlines():
        compact = re.sub(r"\s+", " ", raw_line).strip()
        if not compact or compact.casefold().endswith("| resume"):
            continue
        current_is_structural = bool(
            _section_heading(compact)
            or _is_contact_line(compact)
            or _is_date_line(compact)
            or _is_labeled_line(compact)
            or re.match(r"^[•▪●*-]\s*", compact)
        )
        if prepared:
            previous = prepared[-1]
            previous_text = re.sub(r"^[•▪●*-]\s*", "", previous)
            previous_is_structural = bool(
                _section_heading(previous_text)
                or _is_contact_line(previous_text)
                or _is_date_line(previous_text)
                or _is_labeled_line(previous_text)
            )
            should_join = (
                not current_is_structural
                and not previous_is_structural
                and previous_text[-1:] not in ".:;!?"
            )
            if should_join:
                prepared[-1] = previous + " " + compact
                continue
        prepared.append(compact)
    return prepared


def _job_focus_terms(job_description: str, source_resume: str, limit: int = 28) -> list[str]:
    """Return important JD terms that are supported by the source resume."""
    resume_normalized = clean_text(source_resume)
    terms: list[str] = []
    matched_skills = extract_skills(job_description) & extract_skills(source_resume)
    for skill in sorted(matched_skills):
        matching_aliases = [
            alias
            for alias in SKILL_ALIASES.get(skill, [skill])
            if clean_text(alias) in resume_normalized
        ]
        terms.extend(matching_aliases or [skill])

    tokens = [
        token.strip(".-")
        for token in re.findall(r"[a-z][a-z0-9+#.-]{3,}", clean_text(job_description))
    ]
    counts = Counter(token for token in tokens if token and token not in FOCUS_STOPWORDS)
    terms.extend(
        token
        for token, _ in counts.most_common()
        if re.search(rf"(?<!\w){re.escape(token)}(?!\w)", resume_normalized)
    )

    unique: list[str] = []
    seen: set[str] = set()
    for term in sorted(terms, key=lambda value: (-len(value), value.casefold())):
        normalized = clean_text(term)
        if normalized and normalized not in seen:
            unique.append(term)
            seen.add(normalized)
        if len(unique) >= limit:
            break
    return unique


def build_targeted_questions(
    job_description: str,
    source_resume: str,
    reported_gaps: str = "",
    limit: int = 5,
) -> list[str]:
    """Create quick evidence questions without making another AI request."""
    required_skills, _ = extract_required_skills(job_description)
    missing_skills = sorted(required_skills - extract_skills(source_resume))
    questions = [
        f"Do you have genuine hands-on experience with {skill}?"
        for skill in missing_skills[:3]
    ]
    resume_numbers = re.findall(r"\b\d+(?:\.\d+)?%?\b", source_resume)
    if not resume_numbers:
        questions.append(
            "Provide one verified measurable result for a relevant responsibility or project "
            "(for example time saved, records processed, accuracy, volume or turnaround time)."
        )
    if reported_gaps and "education" in reported_gaps.casefold():
        questions.append("Enter any completed, verifiable qualification or certification relevant to this role that is missing from the resume.")
    questions.append("Describe one job-relevant accomplishment that is completed but not currently written in the resume.")
    return list(dict.fromkeys(questions))[:limit]


def _display_skill(skill: str) -> str:
    """Return familiar résumé casing for a verified skill name."""
    normalized = clean_text(skill)
    preferred = {
        "aws": "AWS", "api": "API", "css": "CSS", "etl": "ETL",
        "html": "HTML", "javascript": "JavaScript", "llm": "LLM",
        "nlp": "NLP", "power bi": "Power BI", "python": "Python",
        "r": "R", "sql": "SQL", "spss": "SPSS", "vba": "VBA",
    }
    return preferred.get(normalized, " ".join(word.capitalize() for word in skill.strip().split()))


def _parse_verified_information(additional_information: str) -> dict[str, list[str]]:
    """Classify user-confirmed facts so they can be merged into real résumé sections."""
    parsed = {"skills": [], "results": [], "accomplishments": [], "qualifications": []}
    for raw_line in additional_information.splitlines():
        line = re.sub(r"^[•\-*]\s*", "", re.sub(r"\s+", " ", raw_line)).strip()
        if len(line) < 2:
            continue
        tagged = re.match(r"^\[(SKILL|RESULT|ACCOMPLISHMENT|QUALIFICATION|EVIDENCE)]\s*(.+)$", line, re.IGNORECASE)
        if tagged:
            category, value = tagged.group(1).casefold(), tagged.group(2).strip()
            target = {
                "skill": "skills", "result": "results", "accomplishment": "accomplishments",
                "qualification": "qualifications", "evidence": "accomplishments",
            }[category]
            parsed[target].append(_display_skill(value) if target == "skills" else value)
            continue
        skill_match = re.match(r"^verified hands-on experience with\s+(.+?)[.]?$", line, re.IGNORECASE)
        if skill_match:
            parsed["skills"].append(_display_skill(skill_match.group(1).strip(" .?")))
        else:
            parsed["accomplishments"].append(line)
    for category, values in parsed.items():
        seen: set[str] = set()
        parsed[category] = [
            value for value in values
            if not (clean_text(value) in seen or seen.add(clean_text(value)))
        ]
    return parsed


def _merge_verified_skills(source_text: str, verified_skills: list[str]) -> str:
    """Add confirmed job skills to the existing skills line without a new evidence section."""
    existing = clean_text(source_text)
    missing = [skill for skill in verified_skills if clean_text(skill) not in existing]
    if not missing:
        return source_text
    separator = " | " if "|" in source_text else ", "
    return source_text.rstrip(" ,;|") + separator + separator.join(missing)


def _evidence_tokens(text: str) -> set[str]:
    return {
        token for token in re.findall(r"[a-z][a-z0-9+#.-]{2,}", clean_text(text))
        if token not in FOCUS_STOPWORDS
    }


def _verified_text_is_job_relevant(text: str, job_description: str, source_resume: str) -> bool:
    """Reject free text that has no clear connection to the job or existing résumé evidence."""
    text_skills = extract_skills(text)
    job_skills = extract_skills(job_description)
    if text_skills & job_skills:
        return True
    text_tokens = _evidence_tokens(text)
    if text_tokens & _evidence_tokens(job_description):
        return True
    return len(text_tokens & _evidence_tokens(source_resume)) >= 2


def _tailorable_lines(source_resume: str, maximum_lines: int = 70) -> list[dict]:
    """Return stable line IDs for summary, skill, experience and project rewriting."""
    current_section = "Profile"
    items: list[dict] = []
    for compact in _prepare_source_lines(source_resume):
        heading = _section_heading(compact)
        if heading:
            current_section = heading
            continue
        if current_section not in TAILORABLE_SECTIONS or _is_contact_line(compact):
            continue
        bullet = bool(re.match(r"^[•▪●*-]\s*", compact))
        text = re.sub(r"^[•▪●*-]\s*", "", compact).strip()
        if not text or _is_date_line(text) and len(text) < 90:
            continue
        if current_section != "Technical Skills" and len(text) < 24 and not bullet:
            continue
        items.append({
            "line_id": f"L{len(items) + 1:03d}",
            "section": current_section,
            "text": text,
            "maximum_characters": max(len(text) + 12, int(len(text) * 1.20)),
        })
        if len(items) >= maximum_lines:
            break
    return items


def _enhance_summary_for_role(
    text: str,
    target_role: str,
    matched_skills: list[str],
    maximum_characters: int,
) -> str:
    """Move verified role evidence forward without adding a new claim."""
    revised = re.sub(r"\s+", " ", text).strip()
    role = re.sub(r"\s+", " ", target_role).strip(" .|-")
    if not role:
        return revised
    display_role = " ".join(word if word.isupper() else word.capitalize() for word in role.split())

    # Replace only an explicitly generic opening. The years, technologies and
    # accomplishments that follow remain exactly the candidate's evidence.
    if re.match(r"^data professional\b", revised, flags=re.IGNORECASE):
        revised = re.sub(
            r"^data professional\b",
            display_role,
            revised,
            count=1,
            flags=re.IGNORECASE,
        )

    # If the source already names the target role, move it to the front of a
    # multi-role summary instead of introducing it as a new qualification.
    with_match = re.search(r"\s+with\s+", revised, flags=re.IGNORECASE)
    prefix_end = with_match.start() if with_match else min(len(revised), 120)
    prefix = revised[:prefix_end]
    role_match = re.search(rf"\b{re.escape(role)}\b", prefix, flags=re.IGNORECASE)
    if role_match and role_match.start() > 0 and not prefix.casefold().startswith(role.casefold()):
        before = prefix[:role_match.start()].strip(" |,/-")
        after = prefix[role_match.end():].strip(" |,/-")
        before = re.sub(r"\b(?:and|or)\s*$", "", before, flags=re.IGNORECASE).strip(" |,/-")
        after = re.sub(r"^(?:and|or)\b", "", after, flags=re.IGNORECASE).strip(" |,/-")
        qualifier = ""
        qualifier_match = re.match(r"^(results[- ]driven|experienced|detail[- ]oriented)\s+", before, re.IGNORECASE)
        if qualifier_match:
            qualifier = qualifier_match.group(0)
            before = before[qualifier_match.end():].strip(" |,/-")
        remainder = " and ".join(part for part in (before, after) if part)
        reordered = f"{qualifier}{display_role}" + (f" and {remainder}" if remainder else "")
        candidate = reordered + revised[prefix_end:]
        if len(candidate) <= maximum_characters:
            revised = candidate

    # Add at most three already-supported requirements only when they are not
    # already visible and the replacement still fits the source text box.
    missing_matches = [
        skill for skill in matched_skills
        if not re.search(rf"(?<!\w){re.escape(skill)}(?!\w)", revised, re.IGNORECASE)
    ][:3]
    if missing_matches:
        strengths = ", ".join(missing_matches[:-1])
        if len(missing_matches) > 1:
            strengths += f" and {missing_matches[-1]}"
        else:
            strengths = missing_matches[0]
        candidate = revised.rstrip(". ") + f". Relevant strengths include {strengths}."
        if len(candidate) <= maximum_characters:
            revised = candidate
    return revised


def _local_tailoring_result(
    job_description: str,
    source_resume: str,
    reported_gaps: str = "",
    target_role: str = "",
    additional_information: str = "",
) -> dict:
    """Fast built-in tailoring using only facts already present in the resume."""
    source_normalized = clean_text(source_resume)
    verified = _parse_verified_information(additional_information)
    verified_skills = [skill for skill in verified["skills"] if clean_text(skill) not in source_normalized]
    written_candidates = {
        "results": [item[:500] for item in verified["results"] if clean_text(item) not in source_normalized],
        "accomplishments": [item[:500] for item in verified["accomplishments"] if clean_text(item) not in source_normalized],
        "qualifications": [item[:500] for item in verified["qualifications"] if clean_text(item) not in source_normalized],
    }
    relevant_written = {
        category: [
            item for item in values
            if _verified_text_is_job_relevant(item, job_description, source_resume)
        ]
        for category, values in written_candidates.items()
    }
    ignored_written = [
        item
        for category, values in written_candidates.items()
        for item in values
        if item not in relevant_written[category]
    ]
    verified_results = relevant_written["results"]
    verified_accomplishments = relevant_written["accomplishments"]
    verified_qualifications = relevant_written["qualifications"]

    required_skills, _ = extract_required_skills(job_description)
    verified_text = verified_skills + verified_results + verified_accomplishments + verified_qualifications
    complete_evidence = source_resume + ("\n" + "\n".join(verified_text) if verified_text else "")
    resume_skills = extract_skills(complete_evidence)
    matches = sorted(required_skills & resume_skills, key=lambda skill: job_description.casefold().find(skill))
    gaps = sorted(required_skills - resume_skills)
    rewrites = []
    source_items = _tailorable_lines(source_resume)
    summary_updated = False
    skill_updated = False
    for item in source_items:
        revised = _enhance_line(item["text"])
        if not summary_updated and item["section"] == "Professional Summary":
            revised = _enhance_summary_for_role(revised, target_role, matches, item["maximum_characters"])
            summary_updated = revised != item["text"]
        if not skill_updated and item["section"] == "Technical Skills" and verified_skills:
            revised = _merge_verified_skills(revised, verified_skills)
            skill_updated = revised != item["text"]
        if revised != item["text"]:
            rewrites.append({
                "line_id": item["line_id"],
                "source_text": item["text"],
                "revised_text": revised,
                "reason": "Prioritized verified job-relevant evidence and strengthened the wording.",
            })
        if len(rewrites) >= 12:
            break

    section_additions: dict[str, list[str]] = {}
    evidence_items = verified_results + verified_accomplishments
    candidate_items = [
        item for item in source_items
        if item["section"] in {"Professional Experience", "Selected Projects"}
    ]
    for evidence in evidence_items:
        evidence_tokens = _evidence_tokens(evidence)
        ranked = sorted(
            candidate_items,
            key=lambda item: len(evidence_tokens & _evidence_tokens(item["text"])),
            reverse=True,
        )
        best = ranked[0] if ranked else None
        overlap = len(evidence_tokens & _evidence_tokens(best["text"])) if best else 0
        if best and overlap >= 2:
            rewrites = [
                rewrite for rewrite in rewrites
                if clean_text(rewrite["source_text"]) != clean_text(best["text"])
            ]
            rewrites.append({
                "line_id": best["line_id"],
                "source_text": best["text"],
                "revised_text": evidence.rstrip(". ") + ".",
                "reason": "Strengthened the existing bullet with a user-verified accomplishment or measurable result.",
            })
            candidate_items.remove(best)
        else:
            section_additions.setdefault("Professional Experience", []).append(evidence.rstrip(". ") + ".")
    if verified_qualifications:
        destination = "Certifications" if "Certifications" in {_section_heading(line) for line in _prepare_source_lines(source_resume)} else "Education"
        section_additions.setdefault(destination, []).extend(
            qualification.rstrip(". ") + "." for qualification in verified_qualifications
        )

    change_summary = [
        "Preserved all source facts, dates, qualifications and sections.",
        "Reworked the summary and weak bullet openings to prioritize verified job-relevant evidence.",
        "Merged confirmed skills into the existing skills section and placed verified answers in the relevant experience, project or education section.",
        "Used only skills and accomplishments supported by the uploaded resume or the candidate's verified answers.",
    ]
    if verified_skills:
        change_summary.insert(0, "Added confirmed skills to Technical Skills: " + ", ".join(verified_skills) + ".")
    if verified_results or verified_accomplishments or verified_qualifications:
        change_summary.insert(
            1 if verified_skills else 0,
            "Applied the candidate's additional verified written information inside the relevant résumé sections.",
        )
    if ignored_written:
        change_summary.append(
            "Skipped additional text that had no clear connection to the job requirements or existing résumé evidence."
        )

    result = {
        "rewrites": rewrites,
        "strong_matches": matches,
        "remaining_gaps": [f"No verified resume evidence for {skill}." for skill in gaps],
        "questions": build_targeted_questions(job_description, source_resume, reported_gaps),
        "verified_skills": verified_skills,
        "section_additions": section_additions,
        "ignored_verified_information": ignored_written,
        "change_summary": change_summary,
        "generation_mode": "built-in evidence-based",
    }
    return result


def _validated_ai_tailoring(
    response: ResumeTailoringResponse,
    source_items: list[dict],
    source_resume: str,
    additional_information: str,
) -> dict:
    """Reject AI edits that add unsupported skills, numbers or oversized text."""
    source_by_id = {item["line_id"]: item for item in source_items}
    evidence = source_resume + "\n" + additional_information
    evidence_skills = extract_skills(evidence)
    evidence_numbers = set(re.findall(r"\b\d+(?:\.\d+)?%?\b", evidence))
    accepted: list[dict] = []
    rejected = 0
    for rewrite in response.rewrites:
        item = source_by_id.get(rewrite.line_id)
        revised = re.sub(r"\s+", " ", rewrite.revised_text).strip()
        if not item or not revised:
            rejected += 1
            continue
        new_skills = extract_skills(revised) - evidence_skills
        new_numbers = set(re.findall(r"\b\d+(?:\.\d+)?%?\b", revised)) - evidence_numbers
        if new_skills or new_numbers or len(revised) > item["maximum_characters"]:
            rejected += 1
            continue
        accepted.append({
            "line_id": rewrite.line_id,
            "source_text": item["text"],
            "revised_text": revised,
            "reason": rewrite.reason.strip(),
        })
    summary = [item.strip() for item in response.change_summary if item.strip()]
    if rejected:
        summary.append(f"Rejected {rejected} suggested edit(s) that were unsupported or too long for the source layout.")
    return {
        "rewrites": accepted,
        "strong_matches": [item.strip() for item in response.strong_matches if item.strip()],
        "remaining_gaps": [item.strip() for item in response.remaining_gaps if item.strip()],
        "questions": [item.strip() for item in response.questions if item.strip()],
        "change_summary": summary,
        "generation_mode": "validated evidence-based",
    }


def tailor_resume_with_ai(
    job_description: str,
    source_resume: str,
    additional_information: str = "",
    reported_gaps: str = "",
    target_role: str = "",
) -> dict:
    """Compatibility wrapper for the built-in, no-key enhancement engine."""
    return _local_tailoring_result(
        job_description,
        source_resume,
        reported_gaps,
        target_role=target_role,
        additional_information=additional_information,
    )


def _enhance_line(line: str) -> str:
    """Apply conservative wording cleanup without introducing new facts."""
    enhanced = re.sub(r"\s+", " ", line).strip()
    changed = False
    replacements = (
        (r"^worked on\s+", "Contributed to "),
        (r"^involved in\s+", "Contributed to "),
        (r"^was involved in\s+", "Contributed to "),
        (r"^responsible for\s+", "Delivered "),
        (r"^utilized\s+", "Used "),
        (r"^develop\s+", "Developed "),
        (r"^analyse\s+", "Analyzed "),
        (r"^analysed\s+", "Analyzed "),
    )
    for pattern, replacement in replacements:
        if re.search(pattern, enhanced, flags=re.IGNORECASE):
            enhanced = re.sub(pattern, replacement, enhanced, count=1, flags=re.IGNORECASE)
            changed = True
            break
    if changed and len(enhanced) >= 35 and enhanced[-1] not in ".:;!?|":
        enhanced += "."
    return enhanced


def build_enhanced_resume(
    candidate_name: str,
    target_role: str,
    job_description: str,
    source_resume: str,
    tailoring_result: dict | None = None,
) -> dict:
    """Keep every source line while cleaning wording and identifying supported focus terms."""
    display_name = _normalize_candidate_name(candidate_name)
    tailoring_result = tailoring_result or {}
    verified_support = "\n".join(
        list(tailoring_result.get("verified_skills", []))
        + [text for values in tailoring_result.get("section_additions", {}).values() for text in values]
    )
    focus_terms = _job_focus_terms(job_description, source_resume + "\n" + verified_support)
    rewrite_map = {
        clean_text(item.get("source_text", "")): item.get("revised_text", "")
        for item in tailoring_result.get("rewrites", [])
        if item.get("source_text") and item.get("revised_text")
    }
    sections: list[dict] = []
    current = {"heading": "Profile", "lines": []}
    sections.append(current)
    contact_lines: list[str] = []

    for compact in _prepare_source_lines(source_resume):
        if clean_text(compact) == clean_text(display_name):
            continue
        heading = _section_heading(compact)
        if heading:
            current = {"heading": heading, "lines": []}
            sections.append(current)
            continue
        if _is_contact_line(compact) and len(sections) == 1:
            contact_lines.append(compact)
            continue
        bullet = bool(re.match(r"^[•▪●*-]\s*", compact))
        cleaned = re.sub(r"^[•▪●*-]\s*", "", compact)
        enhanced = rewrite_map.get(clean_text(cleaned)) or _enhance_line(cleaned)
        relevant = any(clean_text(term) in clean_text(enhanced) for term in focus_terms)
        current["lines"].append({
            "text": enhanced,
            "source": cleaned,
            "bullet": bullet,
            "relevant": relevant,
        })

    sections = [section for section in sections if section["lines"]]
    verified_skills = list(tailoring_result.get("verified_skills", []))
    if verified_skills:
        skill_section = next((item for item in sections if item["heading"] == "Technical Skills"), None)
        if skill_section is None:
            skill_section = {"heading": "Technical Skills", "lines": []}
            sections.append(skill_section)
        existing_skill_text = " ".join(item["text"] for item in skill_section["lines"])
        missing_verified_skills = [
            skill for skill in verified_skills
            if clean_text(skill) not in clean_text(existing_skill_text)
        ]
        if missing_verified_skills and skill_section["lines"]:
            target = max(
                skill_section["lines"],
                key=lambda item: (
                    len(extract_skills(item["text"])),
                    item["text"].count(",") + item["text"].count("|"),
                    len(item["text"]),
                ),
            )
            target["text"] = _merge_verified_skills(target["text"], missing_verified_skills)
            target["relevant"] = True
        elif missing_verified_skills:
            skill_section["lines"].append({
                "text": ", ".join(missing_verified_skills),
                "source": "",
                "bullet": False,
                "relevant": True,
            })
    for destination, additions in tailoring_result.get("section_additions", {}).items():
        section = next((item for item in sections if item["heading"] == destination), None)
        if section is None:
            section = {"heading": destination, "lines": []}
            sections.append(section)
        section["lines"].extend(
            {"text": addition, "source": "", "bullet": destination not in {"Technical Skills"}, "relevant": True}
            for addition in additions
        )
    return {
        "candidate_name": display_name,
        "target_role": target_role.strip(),
        "contact": " | ".join(dict.fromkeys(contact_lines)),
        "focus_terms": focus_terms,
        "sections": sections,
        "tailoring": tailoring_result,
    }


def _split_focus_text(text: str, focus_terms: list[str]) -> list[tuple[str, bool]]:
    terms = [term for term in focus_terms if term]
    if not terms:
        return [(text, False)]
    pattern = re.compile(
        "(" + "|".join(re.escape(term) for term in sorted(terms, key=len, reverse=True)) + ")",
        re.IGNORECASE,
    )
    return [(part, bool(pattern.fullmatch(part))) for part in pattern.split(text) if part]


def _add_docx_text(paragraph, text: str, focus_terms: list[str], base_bold: bool = False) -> None:
    for part, focused in _split_focus_text(text, focus_terms):
        run = paragraph.add_run(part)
        run.bold = base_bold or focused
        run.font.name = "Arial"
        run.font.size = Pt(8.6)
        run.font.color.rgb = RGBColor(25, 25, 25)


def _add_bottom_border(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "333333")
    borders.append(bottom)


def create_enhanced_docx(candidate_name: str, content: dict) -> tuple[str, bytes]:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.38)
    section.left_margin = section.right_margin = Inches(0.55)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.6)
    normal.font.color.rgb = RGBColor(25, 25, 25)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    name_run = name.add_run(content["candidate_name"].upper())
    name_run.bold = True
    name_run.font.name = "Arial"
    name_run.font.size = Pt(18)
    name_run.font.color.rgb = RGBColor(15, 15, 15)

    if content["target_role"]:
        target = document.add_paragraph()
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        target.paragraph_format.space_after = Pt(1)
        target_run = target.add_run(content["target_role"])
        target_run.bold = True
        target_run.font.size = Pt(10)
        target_run.font.color.rgb = RGBColor(55, 55, 55)
    if content["contact"]:
        contact = document.add_paragraph()
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact.paragraph_format.space_after = Pt(5)
        contact_run = contact.add_run(content["contact"])
        contact_run.font.size = Pt(8)
        contact_run.font.color.rgb = RGBColor(70, 70, 70)

    page_break_sources = {clean_text(value) for value in content.get("page_break_after_sources", [])}
    for section_data in content["sections"]:
        if section_data["heading"] != "Profile":
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(4)
            heading.paragraph_format.space_after = Pt(2)
            heading.paragraph_format.keep_with_next = True
            heading_run = heading.add_run(section_data["heading"].upper())
            heading_run.bold = True
            heading_run.font.name = "Arial"
            heading_run.font.size = Pt(9.5)
            heading_run.font.color.rgb = RGBColor(20, 20, 20)
            _add_bottom_border(heading)

        for item in section_data["lines"]:
            text = item["text"]
            is_label = text.casefold().rstrip(":") in {"responsibilities", "project 1", "project 2", "proficient in"}
            is_role_line = bool(
                re.search(r"\b(?:19|20)\d{2}\b", text)
                or ("|" in text and len(text) < 150)
                or is_label
            )
            paragraph = document.add_paragraph(style="List Bullet" if item["bullet"] else None)
            paragraph.paragraph_format.left_indent = Inches(0.16 if item["bullet"] else 0)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12 if item["bullet"] else 0)
            paragraph.paragraph_format.space_after = Pt(0.8)
            paragraph.paragraph_format.keep_together = True
            _add_docx_text(paragraph, text, content["focus_terms"], base_bold=is_role_line)
            if clean_text(item.get("source", "")) in page_break_sources:
                document.add_page_break()

    output = BytesIO()
    document.save(output)
    return f"{_safe_filename(candidate_name)}_enhanced.docx", output.getvalue()


def _pdf_focus_markup(text: str, focus_terms: list[str], base_bold: bool = False) -> str:
    parts = []
    for part, focused in _split_focus_text(text, focus_terms):
        escaped = html_escape(part)
        parts.append(f"<b>{escaped}</b>" if base_bold or focused else escaped)
    return "".join(parts)


def create_enhanced_pdf(candidate_name: str, content: dict) -> tuple[str, bytes]:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=13 * mm,
        bottomMargin=13 * mm,
    )
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "ResumeName", parent=styles["Title"], alignment=TA_CENTER,
        fontName="Helvetica-Bold", fontSize=18, leading=20,
        textColor=colors.HexColor("#111111"), spaceAfter=2,
    )
    target_style = ParagraphStyle(
        "ResumeTarget", parent=styles["Normal"], alignment=TA_CENTER,
        fontName="Helvetica-Bold", fontSize=9.5, leading=11,
        textColor=colors.HexColor("#333333"), spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "ResumeContact", parent=styles["Normal"], alignment=TA_CENTER,
        fontSize=7.8, leading=9.5, textColor=colors.HexColor("#444444"), spaceAfter=5,
    )
    heading_style = ParagraphStyle(
        "ResumeHeading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=10, leading=12, textColor=colors.HexColor("#111111"),
        spaceBefore=6, spaceAfter=3, borderWidth=0, borderPadding=0,
    )
    body_style = ParagraphStyle(
        "ResumeBody", parent=styles["BodyText"], fontSize=8.4, leading=10.8,
        textColor=colors.HexColor("#1A1A1A"), spaceAfter=1.6,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet", parent=body_style, leftIndent=11, firstLineIndent=-7,
        bulletIndent=2, spaceAfter=1.4,
    )

    story = [Paragraph(html_escape(content["candidate_name"].upper()), name_style)]
    if content["target_role"]:
        story.append(Paragraph(html_escape(content["target_role"]), target_style))
    if content["contact"]:
        story.append(Paragraph(html_escape(content["contact"]), contact_style))

    for section_data in content["sections"]:
        if section_data["heading"] != "Profile":
            story.append(Paragraph(html_escape(section_data["heading"].upper()), heading_style))
        for item in section_data["lines"]:
            text = item["text"]
            is_label = text.casefold().rstrip(":") in {"responsibilities", "project 1", "project 2", "proficient in"}
            base_bold = bool(
                re.search(r"\b(?:19|20)\d{2}\b", text)
                or ("|" in text and len(text) < 150)
                or is_label
            )
            markup = _pdf_focus_markup(text, content["focus_terms"], base_bold)
            if item["bullet"]:
                story.append(Paragraph("• " + markup, bullet_style))
            else:
                story.append(Paragraph(markup, body_style))
        story.append(Spacer(1, 1))

    document.build(story)
    return f"{_safe_filename(candidate_name)}_enhanced.pdf", output.getvalue()


def _focus_spans(text: str, focus_terms: list[str]) -> list[tuple[int, int]]:
    """Find and merge every case-insensitive focus-term span in one text pass."""
    spans: list[tuple[int, int]] = []
    for term in sorted({term.strip() for term in focus_terms if len(term.strip()) >= 3}, key=len, reverse=True):
        spans.extend(
            (match.start(), match.end())
            for match in re.finditer(re.escape(term), text, flags=re.IGNORECASE)
        )
    if not spans:
        return []

    merged: list[list[int]] = []
    for start, end in sorted(spans):
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return [(start, end) for start, end in merged]


def _apply_word_focus(document, focus_terms: list[str]) -> None:
    """Bold matched terms without repeatedly searching the complete Word file."""
    content_range = document.Content
    full_text = str(content_range.Text)
    base_position = int(content_range.Start)
    final_position = int(content_range.End)
    for start, end in _focus_spans(full_text, focus_terms):
        absolute_start = min(base_position + start, final_position)
        absolute_end = min(base_position + end, final_position)
        if absolute_end > absolute_start:
            document.Range(absolute_start, absolute_end).Font.Bold = True


def _apply_word_focus_with_find(document, focus_terms: list[str]) -> None:
    """Reliable Word-native fallback used only if direct range emphasis fails."""
    for term in focus_terms:
        if len(term.strip()) < 3:
            continue
        try:
            finder = document.Content.Find
            finder.ClearFormatting()
            finder.Replacement.ClearFormatting()
            finder.Replacement.Font.Bold = True
            finder.Execute(
                FindText=term,
                MatchCase=False,
                MatchWholeWord=False,
                Forward=True,
                Wrap=1,
                Format=False,
                ReplaceWith="^&",
                Replace=2,
            )
        except Exception:
            continue


def _apply_word_focus_safely(document, focus_terms: list[str]) -> None:
    """Keep the Word conversion even when the optimized emphasis route is unsupported."""
    try:
        _apply_word_focus(document, focus_terms)
    except Exception:
        _apply_word_focus_with_find(document, focus_terms)


def _word_page_text_is_blank(text: str) -> bool:
    return not re.sub(r"[\s\r\x07\x0c]+", "", text or "")


def _remove_trailing_blank_word_pages(document) -> None:
    """Delete up to three empty final pages left by Word's PDF conversion."""
    try:
        for _ in range(3):
            page_count = int(document.ComputeStatistics(2))  # wdStatisticPages
            if page_count <= 1:
                return
            page_start = document.GoTo(What=1, Which=1, Count=page_count)  # wdGoToPage / absolute
            last_page = document.Range(page_start.Start, document.Content.End)
            if not _word_page_text_is_blank(str(last_page.Text)):
                return
            last_page.Delete()
    except Exception:
        return


def _remove_blank_pdf_pages(pdf_bytes: bytes) -> bytes:
    """Remove truly empty exported pages without touching pages containing résumé content."""
    try:
        document = fitz.open(stream=pdf_bytes, filetype="pdf")
        if document.page_count <= 1:
            output = document.tobytes(garbage=4, deflate=True)
            document.close()
            return output

        blank_pages: list[int] = []
        for page_number in range(document.page_count):
            page = document[page_number]
            text_blocks = [block for block in page.get_text("blocks") if str(block[4]).strip()]
            visible_text = re.sub(r"\s+", "", " ".join(str(block[4]) for block in text_blocks))
            has_images = bool(page.get_images(full=True))
            has_drawings = bool(page.get_drawings())
            margin_only = bool(text_blocks) and len(visible_text) < 80 and all(
                block[1] <= page.rect.height * 0.12 or block[3] >= page.rect.height * 0.88
                for block in text_blocks
            )
            if (not visible_text or margin_only) and not has_images and not has_drawings:
                blank_pages.append(page_number)

        for page_number in reversed(blank_pages):
            if document.page_count > 1:
                document.delete_page(page_number)
        output = document.tobytes(garbage=4, deflate=True)
        document.close()
        return output
    except Exception:
        return pdf_bytes


def _enhance_original_pdf(
    candidate_name: str,
    source_bytes: bytes,
    focus_terms: list[str],
) -> tuple[str, bytes] | None:
    """Retain original PDF pages and add subtle black emphasis without coloured blocks."""
    try:
        document = fitz.open(stream=source_bytes, filetype="pdf")
        marks_added = 0
        seen_rectangles: set[tuple[int, int, int, int, int]] = set()
        for page_number, page in enumerate(document):
            for term in sorted(focus_terms, key=len, reverse=True):
                if marks_added >= 40:
                    break
                for rectangle in page.search_for(term):
                    key = (
                        page_number,
                        round(rectangle.x0),
                        round(rectangle.y0),
                        round(rectangle.x1),
                        round(rectangle.y1),
                    )
                    if key in seen_rectangles:
                        continue
                    seen_rectangles.add(key)
                    underline_y = min(rectangle.y1 + 0.35, page.rect.y1 - 0.5)
                    page.draw_line(
                        (rectangle.x0, underline_y),
                        (rectangle.x1, underline_y),
                        color=(0, 0, 0),
                        width=0.45,
                        overlay=True,
                    )
                    marks_added += 1
                    if marks_added >= 40:
                        break
        output = document.tobytes(garbage=4, deflate=True)
        document.close()
        return f"{_safe_filename(candidate_name)}_enhanced.pdf", _remove_blank_pdf_pages(output)
    except Exception:
        return None


def _pdf_page_break_sources(source_bytes: bytes) -> list[str]:
    """Identify the final content line on every source page except the last."""
    document = fitz.open(stream=source_bytes, filetype="pdf")
    break_sources: list[str] = []
    for page in list(document)[:-1]:
        lines = _prepare_source_lines(page.get_text("text"))
        candidates = [
            re.sub(r"^[•▪●*-]\s*", "", line).strip()
            for line in lines
            if line.strip() and not _is_contact_line(line) and not _section_heading(line)
        ]
        if candidates:
            break_sources.append(candidates[-1])
    document.close()
    return break_sources


def _enhance_text_pdf_in_place(
    candidate_name: str,
    source_bytes: bytes,
    tailoring_result: dict,
) -> tuple[str, bytes, list[str]]:
    """Apply bounded text rewrites on the original searchable PDF pages."""
    document = fitz.open(stream=source_bytes, filetype="pdf")
    operations: dict[int, list[dict]] = {}
    warnings: list[str] = []
    for rewrite in tailoring_result.get("rewrites", []):
        source_text = re.sub(r"\s+", " ", rewrite.get("source_text", "")).strip()
        revised_text = re.sub(r"\s+", " ", rewrite.get("revised_text", "")).strip()
        punctuation_free_source = re.sub(r"[^a-z0-9]+", "", source_text.casefold())
        punctuation_free_revised = re.sub(r"[^a-z0-9]+", "", revised_text.casefold())
        if (
            not source_text
            or not revised_text
            or source_text == revised_text
            or punctuation_free_source == punctuation_free_revised
        ):
            continue
        located = False
        for page_number, page in enumerate(document):
            rectangles = page.search_for(source_text)
            if not rectangles:
                continue
            union = fitz.Rect(rectangles[0])
            for rectangle in rectangles[1:]:
                union |= rectangle
            union.x0 = max(page.rect.x0, union.x0 - 1.5)
            union.y0 = max(page.rect.y0, union.y0 - 1.0)
            union.x1 = min(
                page.rect.x1 - 18.0,
                max(union.x1 + 2.0, union.x0 + (union.width * 1.08)),
            )
            union.y1 = min(page.rect.y1, union.y1 + 1.5)

            sizes = []
            matching_spans = []
            for block in page.get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        if fitz.Rect(span["bbox"]).intersects(union):
                            sizes.append(float(span.get("size", 9.0)))
                            matching_spans.append(span)
            font_size = sorted(sizes)[len(sizes) // 2] if sizes else 9.0
            source_font = (matching_spans[0].get("font", "") if matching_spans else "").casefold()
            source_flags = int(matching_spans[0].get("flags", 0)) if matching_spans else 0
            is_serif = any(
                token in source_font
                for token in ("times", "cambria", "georgia", "garamond", "serif")
            )
            is_bold = "bold" in source_font or bool(source_flags & 16)
            is_italic = any(token in source_font for token in ("italic", "oblique")) or bool(source_flags & 2)
            if is_serif:
                font_name = "tibi" if is_bold and is_italic else "tibo" if is_bold else "tiit" if is_italic else "tiro"
            else:
                font_name = "hebi" if is_bold and is_italic else "hebo" if is_bold else "heit" if is_italic else "helv"
            source_color = int(matching_spans[0].get("color", 0)) if matching_spans else 0
            color = (
                ((source_color >> 16) & 255) / 255,
                ((source_color >> 8) & 255) / 255,
                (source_color & 255) / 255,
            )
            operations.setdefault(page_number, []).append({
                "rect": union,
                "text": revised_text,
                "source_length": len(source_text),
                "font_size": font_size,
                "font_family": "serif" if is_serif else "sans-serif",
                "font_weight": "bold" if is_bold else "normal",
                "font_style": "italic" if is_italic else "normal",
                "color": color,
                "scale_low": 0.72 if len(revised_text) > len(source_text) else 0.82,
            })
            located = True
            break
        if not located:
            warnings.append(f"Could not safely place one rewrite beginning: {source_text[:55]}")

    for page_number, page_operations in operations.items():
        page = document[page_number]
        # A model may return both a full multi-line rewrite and a rewrite of a
        # sentence inside it. Applying both creates doubled/overlapping text.
        # Keep the most complete rewrite for each occupied region.
        non_overlapping: list[dict] = []
        for operation in sorted(
            page_operations,
            key=lambda item: (item["source_length"], item["rect"].get_area()),
            reverse=True,
        ):
            overlaps_existing = False
            for selected in non_overlapping:
                intersection = operation["rect"] & selected["rect"]
                smaller_area = min(
                    operation["rect"].get_area(),
                    selected["rect"].get_area(),
                )
                if intersection.get_area() / max(smaller_area, 1.0) >= 0.20:
                    overlaps_existing = True
                    break
            if overlaps_existing:
                warnings.append(
                    f"Skipped one overlapping rewrite on page {page_number + 1} to protect the original layout."
                )
                continue
            non_overlapping.append(operation)
        page_operations = sorted(
            non_overlapping,
            key=lambda item: (item["rect"].y0, item["rect"].x0),
        )
        safe_operations = []
        for operation in page_operations:
            red, green, blue = operation["color"]
            color_hex = f"#{round(red * 255):02x}{round(green * 255):02x}{round(blue * 255):02x}"
            operation["html"] = f"<p>{html_escape(operation['text'])}</p>"
            operation["css"] = (
                "p { margin: 0; padding: 0; "
                f"font-family: {operation['font_family']}; "
                f"font-size: {operation['font_size']:.2f}pt; line-height: 1.0; "
                f"font-weight: {operation['font_weight']}; font-style: {operation['font_style']}; "
                f"color: {color_hex}; }}"
            )
            probe = fitz.open()
            probe_page = probe.new_page(
                width=max(operation["rect"].width, 1),
                height=max(operation["rect"].height, 1),
            )
            spare_height, scale = probe_page.insert_htmlbox(
                probe_page.rect,
                operation["html"],
                css=operation["css"],
                scale_low=operation["scale_low"],
            )
            probe.close()
            if spare_height >= 0 and scale >= operation["scale_low"]:
                safe_operations.append(operation)
            else:
                warnings.append(
                    f"A revised line did not fit safely on page {page_number + 1}; the original PDF line was preserved."
                )

        for operation in safe_operations:
            # Remove only the replaced text. A transparent redaction retains
            # coloured bars, shapes and other design elements behind the text.
            page.add_redact_annot(operation["rect"], fill=None)
        if not safe_operations:
            continue
        page.apply_redactions(images=0, graphics=0)
        for operation in safe_operations:
                page.insert_htmlbox(
                operation["rect"],
                operation["html"],
                css=operation["css"],
                scale_low=operation["scale_low"],
                overlay=True,
            )

    output = document.tobytes(garbage=4, deflate=True)
    document.close()
    return f"{_safe_filename(candidate_name)}_enhanced.pdf", _remove_blank_pdf_pages(output), warnings


def _create_from_source_with_microsoft_word(
    candidate_name: str,
    source_name: str,
    source_bytes: bytes,
    focus_terms: list[str],
    word_application=None,
    export_pdf: bool = True,
) -> dict | None:
    """Enhance the uploaded PDF/DOCX through Word while retaining its layout."""
    if platform.system() != "Windows" or Path(source_name).suffix.casefold() not in {".pdf", ".docx"}:
        return None
    owned_session = None
    word = word_application
    document = None
    try:
        if word is None:
            owned_session = _start_microsoft_word()
            if owned_session is None:
                return None
            word = owned_session[1]
        with tempfile.TemporaryDirectory(prefix="matchmind_source_") as directory:
            working_directory = Path(directory)
            source_path = working_directory / ("source" + Path(source_name).suffix.casefold())
            docx_path = working_directory / "enhanced_resume.docx"
            pdf_path = working_directory / "enhanced_resume.pdf"
            source_path.write_bytes(source_bytes)

            document = word.Documents.Open(
                str(source_path.resolve()),
                ConfirmConversions=False,
                ReadOnly=False,
                AddToRecentFiles=False,
            )

            # Strengthen weak openings without adding any unsupported fact.
            for old, new in (
                ("Worked on ", "Contributed to "),
                ("Involved in ", "Contributed to "),
                ("Responsible for ", "Delivered "),
                ("Utilized ", "Used "),
            ):
                finder = document.Content.Find
                finder.ClearFormatting()
                finder.Replacement.ClearFormatting()
                finder.Execute(
                    FindText=old,
                    MatchCase=False,
                    MatchWholeWord=False,
                    Forward=True,
                    Wrap=1,
                    Format=False,
                    ReplaceWith=new,
                    Replace=2,
                )

            # One in-memory scan replaces a full Word search for every term.
            _apply_word_focus_safely(document, focus_terms)
            _remove_trailing_blank_word_pages(document)

            document.SaveAs2(str(docx_path.resolve()), FileFormat=16)
            result = {
                "docx_name": f"{_safe_filename(candidate_name)}_enhanced.docx",
                "docx_bytes": docx_path.read_bytes(),
                "matching_layout": export_pdf,
                "source_layout_preserved": True,
                "layout_note": (
                    "The uploaded DOCX layout was retained by Microsoft Word."
                    if source_path.suffix == ".docx"
                    else "Microsoft Word converted the uploaded PDF and retained its layout as closely as editable Word allows."
                ),
            }
            if export_pdf:
                document.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
                result.update({
                    "pdf_name": f"{_safe_filename(candidate_name)}_enhanced.pdf",
                    "pdf_bytes": _remove_blank_pdf_pages(pdf_path.read_bytes()),
                })
            return result
    except Exception:
        return None
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if owned_session is not None:
            _stop_microsoft_word(owned_session)


def _iter_docx_paragraphs(document: Document):
    """Yield paragraphs in body, tables, headers and footers."""
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for paragraph in section.footer.paragraphs:
            yield paragraph


def _move_body_footer_lines_to_page_footer(document: Document) -> None:
    """Prevent PDF-conversion footer text from creating otherwise blank pages."""
    body_footers = [
        paragraph for paragraph in document.paragraphs
        if re.fullmatch(r".+?\s*\|\s*resume", paragraph.text.strip(), re.IGNORECASE)
    ]
    if not body_footers:
        return
    for index, section in enumerate(document.sections):
        if index:
            section.footer.is_linked_to_previous = False
        source = body_footers[min(index, len(body_footers) - 1)]
        footer = section.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(source.text.strip())
        run.font.name = "Arial"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(80, 80, 80)
    for paragraph in body_footers:
        parent = paragraph._p.getparent()
        if parent is not None:
            parent.remove(paragraph._p)


def _set_paragraph_text_preserving_style(paragraph, text: str, focus_terms: list[str]) -> None:
    """Rewrite paragraph text while retaining every original run's formatting.

    Word frequently splits one résumé line into several runs with different font,
    size, colour, bold, italic, underline, or language settings. Reusing only the
    first run turns designed résumés into plain-looking documents. This mapping
    keeps unchanged characters on their original runs and assigns newly written
    characters the nearest source style.
    """
    del focus_terms  # Job focus comes from truthful wording, not new colour/bold markup.
    original_runs = list(paragraph.runs)
    if not original_runs:
        paragraph.add_run(text)
        return

    source_text = "".join(run.text for run in original_runs)
    run_properties = [
        deepcopy(run._r.rPr) if run._r.rPr is not None else None
        for run in original_runs
    ]
    source_styles: list[int] = []
    for index, run in enumerate(original_runs):
        source_styles.extend([index] * len(run.text))
    if not source_styles:
        source_styles = [0]

    revised_styles: list[int] = []
    matcher = SequenceMatcher(None, source_text, text, autojunk=False)
    for tag, source_start, source_end, revised_start, revised_end in matcher.get_opcodes():
        revised_length = revised_end - revised_start
        if tag == "equal":
            revised_styles.extend(source_styles[source_start:source_end])
            continue
        if revised_length <= 0:
            continue
        if source_start < len(source_styles):
            nearest_style = source_styles[source_start]
        elif source_start:
            nearest_style = source_styles[min(source_start - 1, len(source_styles) - 1)]
        else:
            nearest_style = source_styles[0]
        revised_styles.extend([nearest_style] * revised_length)

    if len(revised_styles) != len(text):
        revised_styles = [source_styles[0]] * len(text)

    for run in original_runs:
        paragraph._p.remove(run._r)

    start = 0
    while start < len(text):
        style_index = revised_styles[start]
        end = start + 1
        while end < len(text) and revised_styles[end] == style_index:
            end += 1
        run = paragraph.add_run(text[start:end])
        properties = run_properties[style_index]
        if properties is not None:
            run._r.insert(0, deepcopy(properties))
        start = end

    if not text:
        run = paragraph.add_run("")
        if run_properties[0] is not None:
            run._r.insert(0, deepcopy(run_properties[0]))


def _find_rewrite_for_paragraph(paragraph_text: str, rewrite_map: dict[str, str]) -> str | None:
    normalized = clean_text(re.sub(r"^[•▪●*-]\s*", "", paragraph_text))
    if normalized in rewrite_map:
        return rewrite_map[normalized]
    for source, revised in rewrite_map.items():
        if len(source) >= 10 and source in normalized:
            source_words = [re.escape(word) for word in source.split()]
            replaced = re.sub(
                r"[\s,;|/•·-]+".join(source_words),
                revised,
                paragraph_text,
                count=1,
                flags=re.IGNORECASE,
            )
            if replaced != paragraph_text:
                return replaced
        length_ratio = len(normalized) / max(len(source), 1)
        if len(source) >= 28 and 0.78 <= length_ratio <= 1.28 and (source in normalized or normalized in source):
            return revised
    return None


def _insert_section_additions_docx(document: Document, section_additions: dict[str, list[str]]) -> None:
    """Place new verified facts inside their normal résumé sections without a separate evidence block."""
    for destination, additions in section_additions.items():
        for addition in additions:
            if clean_text(addition) in clean_text("\n".join(paragraph.text for paragraph in document.paragraphs)):
                continue
            paragraphs = document.paragraphs
            heading_index = next(
                (index for index, paragraph in enumerate(paragraphs) if _section_heading(paragraph.text) == destination),
                None,
            )
            if heading_index is None:
                heading = document.add_paragraph()
                heading_run = heading.add_run(destination.upper())
                heading_run.bold = True
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(addition)
                continue
            anchor_index = next(
                (
                    index for index in range(heading_index + 1, len(paragraphs))
                    if _section_heading(paragraphs[index].text)
                ),
                None,
            )
            section_body = paragraphs[heading_index + 1:anchor_index]
            template = next((paragraph for paragraph in reversed(section_body) if paragraph.text.strip()), None)
            anchor = paragraphs[anchor_index] if anchor_index is not None else None
            paragraph = anchor.insert_paragraph_before() if anchor is not None else document.add_paragraph()
            if template is not None and template._p.pPr is not None:
                paragraph._p.insert(0, deepcopy(template._p.pPr))
            elif destination in {"Professional Experience", "Selected Projects", "Achievements"}:
                paragraph.style = "List Bullet"
            run = paragraph.add_run(addition)
            if template is not None and template.runs and template.runs[0]._r.rPr is not None:
                run._r.insert(0, deepcopy(template.runs[0]._r.rPr))


def _insert_verified_skills_docx(
    document: Document,
    verified_skills: list[str],
    focus_terms: list[str],
) -> int:
    """Guarantee confirmed skills are visible even when the source DOCX uses tables."""
    paragraphs = [paragraph for paragraph in _iter_docx_paragraphs(document) if paragraph.text.strip()]
    document_text = clean_text("\n".join(paragraph.text for paragraph in paragraphs))
    missing = [skill for skill in verified_skills if clean_text(skill) not in document_text]
    if not missing:
        return 0

    heading_index = next(
        (index for index, paragraph in enumerate(paragraphs) if _section_heading(paragraph.text) == "Technical Skills"),
        None,
    )
    candidates = []
    if heading_index is not None:
        for paragraph in paragraphs[heading_index + 1:]:
            if _section_heading(paragraph.text):
                break
            candidates.append(paragraph)
    if not candidates:
        candidates = [
            paragraph for paragraph in paragraphs
            if extract_skills(paragraph.text) and not _section_heading(paragraph.text)
        ]

    if candidates:
        target = max(
            candidates,
            key=lambda paragraph: (
                len(extract_skills(paragraph.text)),
                paragraph.text.count(",") + paragraph.text.count("|"),
                len(paragraph.text),
            ),
        )
        revised = _merge_verified_skills(target.text, missing)
        _set_paragraph_text_preserving_style(target, revised, focus_terms)
    else:
        heading = document.add_paragraph()
        heading_run = heading.add_run("TECHNICAL SKILLS")
        heading_run.bold = True
        paragraph = document.add_paragraph()
        paragraph.add_run(", ".join(missing))
    return len(missing)


def _enhance_source_docx(
    candidate_name: str,
    source_bytes: bytes,
    focus_terms: list[str],
    tailoring_result: dict | None = None,
) -> tuple[str, bytes] | None:
    """Preserve a source DOCX package while applying evidence-checked rewrites."""
    try:
        document = Document(BytesIO(source_bytes))
        _move_body_footer_lines_to_page_footer(document)
        rewrite_map = {
            clean_text(item.get("source_text", "")): item.get("revised_text", "")
            for item in (tailoring_result or {}).get("rewrites", [])
            if item.get("source_text") and item.get("revised_text")
        }
        for paragraph in _iter_docx_paragraphs(document):
            revised = _find_rewrite_for_paragraph(paragraph.text, rewrite_map)
            if revised:
                _set_paragraph_text_preserving_style(paragraph, revised, focus_terms)
                continue
            for run in paragraph.runs:
                updated = run.text
                for pattern, replacement in (
                    (r"\bworked on\b", "Contributed to"),
                    (r"\binvolved in\b", "Contributed to"),
                    (r"\bresponsible for\b", "Delivered"),
                    (r"\butilized\b", "Used"),
                ):
                    updated = re.sub(pattern, replacement, updated, flags=re.IGNORECASE)
                if updated != run.text:
                    run.text = updated
        _insert_section_additions_docx(
            document,
            dict((tailoring_result or {}).get("section_additions", {})),
        )
        _insert_verified_skills_docx(
            document,
            list((tailoring_result or {}).get("verified_skills", [])),
            focus_terms,
        )
        output = BytesIO()
        document.save(output)
        return f"{_safe_filename(candidate_name)}_enhanced.docx", output.getvalue()
    except Exception:
        return None


def _pdf_text_page_count(source_bytes: bytes) -> tuple[bool, int]:
    """Return whether a PDF is text based and its nonblank source page count."""
    document = fitz.open(stream=source_bytes, filetype="pdf")
    text_length = 0
    nonblank_pages = 0
    for page in document:
        text = re.sub(r"\s+", "", page.get_text("text"))
        text_length += len(text)
        if text or page.get_images(full=True) or page.get_drawings():
            nonblank_pages += 1
    document.close()
    return text_length >= 30, max(nonblank_pages, 1)


def _convert_pdf_to_docx_fast(source_bytes: bytes) -> bytes | None:
    """Convert a text PDF structurally without opening Microsoft Word."""
    try:
        from pdf2docx import Converter

        with tempfile.TemporaryDirectory(prefix="matchmind_pdf2docx_") as directory:
            source_path = Path(directory) / "source.pdf"
            output_path = Path(directory) / "source.docx"
            source_path.write_bytes(source_bytes)
            converter = Converter(str(source_path))
            try:
                converter.convert(str(output_path), start=0, end=None)
            finally:
                converter.close()
            if not output_path.exists():
                return None
            document = Document(str(output_path))
            source_document = fitz.open(stream=source_bytes, filetype="pdf")
            first_page = source_document[0].rect
            source_document.close()
            for section in document.sections:
                section.page_width = Pt(first_page.width)
                section.page_height = Pt(first_page.height)
            adjusted = BytesIO()
            document.save(adjusted)
            return adjusted.getvalue()
    except Exception:
        return None


_VML_NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)


def _pdf_colour_hex(colour: int | tuple | list | None, default: str = "000000") -> str:
    if isinstance(colour, int):
        return f"{(colour >> 16) & 255:02X}{(colour >> 8) & 255:02X}{colour & 255:02X}"
    if isinstance(colour, (tuple, list)) and len(colour) >= 3:
        return "".join(f"{max(0, min(255, round(float(value) * 255))):02X}" for value in colour[:3])
    return default


def _pdf_font_attributes(font_name: str, flags: int) -> tuple[str, bool, bool]:
    cleaned = re.sub(r"^[A-Z]{6}\+", "", str(font_name or "Arial"))
    lowered = cleaned.casefold()
    bold = "bold" in lowered or bool(flags & 16)
    italic = any(value in lowered for value in ("italic", "oblique")) or bool(flags & 2)
    if "inriaserif" in lowered or "charissil" in lowered:
        cleaned = "Nimbus Roman"
    elif "times" in lowered:
        cleaned = "Times New Roman"
    elif "helvetica" in lowered:
        cleaned = "Arial"
    else:
        cleaned = re.sub(r"[-,]?(bold|italic|oblique|regular)", "", cleaned, flags=re.IGNORECASE).strip(" -,")
    if not cleaned:
        cleaned = "Arial"
    return cleaned, bold, italic


def _append_positioned_text_box(
    paragraph,
    *,
    shape_id: str,
    text: str,
    bbox,
    page_width: float,
    font_name: str,
    font_size: float,
    colour: str,
    bold: bool,
    italic: bool,
) -> None:
    x0, y0, x1, y1 = (float(value) for value in bbox)
    y0 = max(0.0, y0 - max(1.0, font_size * 0.15))
    # Transparent text boxes may safely overlap. Extending each one to the
    # right page edge prevents Word/LibreOffice from wrapping a PDF text span
    # merely because its substitute font has slightly different metrics.
    width = max((page_width * 1.50) - x0, x1 - x0 + 6.0, 2.0)
    height = max(y1 - y0 + 4.0, font_size * 2.1)
    bold_xml = "<w:b/>" if bold else ""
    italic_xml = "<w:i/>" if italic else ""
    xml = f"""
    <w:r {_VML_NAMESPACES}>
      <w:pict>
        <v:shape id="{shape_id}" type="#_x0000_t202" stroked="f" filled="f"
          style="position:absolute;margin-left:{x0:.2f}pt;margin-top:{y0:.2f}pt;width:{width:.2f}pt;height:{height:.2f}pt;z-index:3;mso-position-horizontal-relative:page;mso-position-vertical-relative:page">
          <v:textbox inset="0,0,0,0">
            <w:txbxContent>
              <w:p>
                <w:pPr><w:spacing w:before="0" w:after="0" w:line="{max(20, round(font_size * 20))}" w:lineRule="exact"/></w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="{html_escape(font_name, quote=True)}" w:hAnsi="{html_escape(font_name, quote=True)}"/>
                    <w:color w:val="{colour}"/><w:sz w:val="{max(2, round(font_size * 2))}"/>
                    {bold_xml}{italic_xml}
                  </w:rPr>
                  <w:t xml:space="preserve">{html_escape(text)}</w:t>
                </w:r>
              </w:p>
            </w:txbxContent>
          </v:textbox>
        </v:shape>
      </w:pict>
    </w:r>
    """
    paragraph._p.append(parse_xml(xml))


def _append_positioned_text_group(
    paragraph,
    *,
    shape_id: str,
    spans: list[dict],
    page_width: float,
) -> None:
    """Place adjacent PDF spans in one Word textbox to prevent run overlap."""
    if not spans:
        return
    x0 = min(float(span["bbox"][0]) for span in spans)
    y0 = min(float(span["bbox"][1]) for span in spans)
    y1 = max(float(span["bbox"][3]) for span in spans)
    maximum_size = max(float(span.get("size", 9.0)) for span in spans)
    source_line_width = max(float(span["bbox"][2]) for span in spans) - x0
    available_line_width = max(page_width - x0, 1.0)
    line_fill_ratio = source_line_width / available_line_width
    serif_substitute_scale = 0.86 if line_fill_ratio >= 0.88 else 0.90
    y0 = max(0.0, y0 - max(1.0, maximum_size * 0.15))
    # PDF lines are already wrapped at their intended visual boundary. Give
    # the Word textbox extra off-page width so a substitute font cannot wrap
    # the same line again and collide with the following positioned line.
    width = max((page_width * 1.50) - x0, 2.0)
    height = max(y1 - y0 + 4.0, maximum_size * 2.1)

    runs_xml: list[str] = []
    previous_x1 = None
    previous_text = ""
    for span in spans:
        span_x0 = float(span["bbox"][0])
        raw_text = str(span.get("text", ""))
        if (
            previous_x1 is not None
            and span_x0 - previous_x1 > 1.2
            and previous_text
            and not previous_text.endswith((" ", "\t"))
            and raw_text
            and not raw_text.startswith((" ", "\t", ",", ".", ":", ";", ")", "]"))
        ):
            raw_text = " " + raw_text
        font_name, bold, italic = _pdf_font_attributes(
            span.get("font", ""),
            int(span.get("flags", 0)),
        )
        bold_xml = "<w:b/>" if bold else ""
        italic_xml = "<w:i/>" if italic else ""
        font_size = float(span.get("size", 9.0))
        if any(
            family in str(span.get("font", "")).casefold()
            for family in ("inriaserif", "charissil")
        ):
            font_size *= serif_substitute_scale
        runs_xml.append(
            f"""
            <w:r>
              <w:rPr>
                <w:rFonts w:ascii="{html_escape(font_name, quote=True)}" w:hAnsi="{html_escape(font_name, quote=True)}"/>
                <w:color w:val="{_pdf_colour_hex(span.get('color', 0))}"/>
                <w:sz w:val="{max(2, round(font_size * 2))}"/>
                {bold_xml}{italic_xml}
              </w:rPr>
              <w:t xml:space="preserve">{html_escape(raw_text)}</w:t>
            </w:r>
            """
        )
        previous_x1 = float(span["bbox"][2])
        previous_text = raw_text

    xml = f"""
    <w:r {_VML_NAMESPACES}>
      <w:pict>
        <v:shape id="{shape_id}" type="#_x0000_t202" stroked="f" filled="f"
          style="position:absolute;margin-left:{x0:.2f}pt;margin-top:{y0:.2f}pt;width:{width:.2f}pt;height:{height:.2f}pt;z-index:3;mso-position-horizontal-relative:page;mso-position-vertical-relative:page">
          <v:textbox inset="0,0,0,0">
            <w:txbxContent>
              <w:p>
                <w:pPr><w:spacing w:before="0" w:after="0" w:line="{max(20, round(maximum_size * 20))}" w:lineRule="exact"/></w:pPr>
                {''.join(runs_xml)}
              </w:p>
            </w:txbxContent>
          </v:textbox>
        </v:shape>
      </w:pict>
    </w:r>
    """
    paragraph._p.append(parse_xml(xml))


def _append_positioned_drawing(paragraph, *, shape_id: str, drawing: dict) -> None:
    rectangle = drawing.get("rect")
    if rectangle is None:
        return
    x0, y0, x1, y1 = (float(value) for value in rectangle)
    width = max(x1 - x0, 0.7)
    height = max(y1 - y0, 0.7)
    stroke = drawing.get("color")
    fill = drawing.get("fill")
    stroke_hex = _pdf_colour_hex(stroke)
    fill_hex = _pdf_colour_hex(fill, "FFFFFF")
    stroke_width = max(float(drawing.get("width") or 0.5), 0.35)
    filled = "t" if fill is not None else "f"
    stroked = "t" if stroke is not None else "f"
    xml = f"""
    <w:r {_VML_NAMESPACES}>
      <w:pict>
        <v:rect id="{shape_id}" filled="{filled}" fillcolor="#{fill_hex}"
          stroked="{stroked}" strokecolor="#{stroke_hex}" strokeweight="{stroke_width:.2f}pt"
          style="position:absolute;margin-left:{x0:.2f}pt;margin-top:{y0:.2f}pt;width:{width:.2f}pt;height:{height:.2f}pt;z-index:1;mso-position-horizontal-relative:page;mso-position-vertical-relative:page"/>
      </w:pict>
    </w:r>
    """
    paragraph._p.append(parse_xml(xml))


def _append_positioned_image(document: Document, paragraph, *, shape_id: str, block: dict) -> None:
    image_bytes = block.get("image")
    bbox = block.get("bbox")
    if not image_bytes or bbox is None:
        return
    try:
        relationship_id, _ = document.part.get_or_add_image(BytesIO(image_bytes))
    except Exception:
        return
    x0, y0, x1, y1 = (float(value) for value in bbox)
    width = max(x1 - x0, 1.0)
    height = max(y1 - y0, 1.0)
    xml = f"""
    <w:r {_VML_NAMESPACES}>
      <w:pict>
        <v:shape id="{shape_id}" type="#_x0000_t75" stroked="f"
          style="position:absolute;margin-left:{x0:.2f}pt;margin-top:{y0:.2f}pt;width:{width:.2f}pt;height:{height:.2f}pt;z-index:2;mso-position-horizontal-relative:page;mso-position-vertical-relative:page">
          <v:imagedata r:id="{relationship_id}" o:title=""/>
        </v:shape>
      </w:pict>
    </w:r>
    """
    paragraph._p.append(parse_xml(xml))


def _append_page_background(
    document: Document,
    paragraph,
    *,
    shape_id: str,
    png_bytes: bytes,
    page_width: float,
    page_height: float,
) -> None:
    relationship_id, _ = document.part.get_or_add_image(BytesIO(png_bytes))
    width_emu = max(int(page_width * 12700), 1)
    height_emu = max(int(page_height * 12700), 1)
    identifier_match = re.search(r"(\d+)$", shape_id)
    drawing_id = int(identifier_match.group(1)) if identifier_match else 1
    xml = f"""
    <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
         xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
         xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
         xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
         xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">
      <w:drawing>
        <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"
                   relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
          <wp:extent cx="{width_emu}" cy="{height_emu}"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          <wp:docPr id="{drawing_id}" name="{shape_id}"/>
          <wp:cNvGraphicFramePr>
            <a:graphicFrameLocks noChangeAspect="1"/>
          </wp:cNvGraphicFramePr>
          <a:graphic>
            <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
              <pic:pic>
                <pic:nvPicPr>
                  <pic:cNvPr id="{drawing_id}" name="{shape_id}.png"/>
                  <pic:cNvPicPr><a:picLocks noChangeAspect="1" noChangeArrowheads="1"/></pic:cNvPicPr>
                </pic:nvPicPr>
                <pic:blipFill>
                  <a:blip r:embed="{relationship_id}" cstate="print"/>
                  <a:srcRect/>
                  <a:stretch><a:fillRect/></a:stretch>
                </pic:blipFill>
                <pic:spPr bwMode="auto">
                  <a:xfrm><a:off x="0" y="0"/><a:ext cx="{width_emu}" cy="{height_emu}"/></a:xfrm>
                  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                  <a:noFill/>
                  <a:ln><a:noFill/></a:ln>
                </pic:spPr>
              </pic:pic>
            </a:graphicData>
          </a:graphic>
        </wp:anchor>
      </w:drawing>
    </w:r>
    """
    paragraph._p.append(parse_xml(xml))


def _render_graphics_only_page(source: fitz.Document, page_number: int) -> bytes:
    """Render non-text design elements while keeping DOCX text fully editable."""
    graphics_document = fitz.open()
    graphics_document.insert_pdf(source, from_page=page_number, to_page=page_number)
    page = graphics_document[0]
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                rectangle = fitz.Rect(span["bbox"])
                rectangle.x0 = max(page.rect.x0, rectangle.x0 - 0.5)
                rectangle.y0 = max(page.rect.y0, rectangle.y0 - 0.5)
                rectangle.x1 = min(page.rect.x1, rectangle.x1 + 0.5)
                rectangle.y1 = min(page.rect.y1, rectangle.y1 + 0.5)
                page.add_redact_annot(rectangle, fill=None)
    page.apply_redactions(images=0, graphics=0)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
    payload = pixmap.tobytes("png")
    graphics_document.close()
    return payload


def _convert_pdf_to_positioned_docx(pdf_bytes: bytes) -> bytes:
    """Recreate PDF pages as editable, selectable, positioned Word content.

    Unlike paragraph-flow PDF converters, this retains columns, colours, text
    placement, lines, boxes and images without flattening the résumé into a page
    screenshot.
    """
    source = fitz.open(stream=pdf_bytes, filetype="pdf")
    if source.page_count == 0:
        source.close()
        raise ValueError("The PDF contains no pages.")

    document = Document()
    shape_counter = 0
    for page_number, page in enumerate(source):
        section = document.sections[-1]
        section.page_width = Pt(page.rect.width)
        section.page_height = Pt(page.rect.height)
        section.top_margin = Pt(0)
        section.bottom_margin = Pt(0)
        section.left_margin = Pt(0)
        section.right_margin = Pt(0)
        section.header_distance = Pt(0)
        section.footer_distance = Pt(0)

        anchor = document.add_paragraph()
        anchor.paragraph_format.space_before = Pt(0)
        anchor.paragraph_format.space_after = Pt(0)
        anchor.paragraph_format.line_spacing = Pt(1)

        shape_counter += 1
        _append_page_background(
            document,
            anchor,
            shape_id=f"matchmind_background_{shape_counter}",
            png_bytes=_render_graphics_only_page(source, page_number),
            page_width=float(page.rect.width),
            page_height=float(page.rect.height),
        )

        page_dict = page.get_text("dict")
        for block in page_dict.get("blocks", []):
            if block.get("type") == 1:
                continue
            for line in block.get("lines", []):
                spans = [span for span in line.get("spans", []) if span.get("text", "")]
                if not spans:
                    continue
                spans.sort(key=lambda span: float(span["bbox"][0]))
                groups: list[list[dict]] = []
                for span in spans:
                    if not groups:
                        groups.append([span])
                        continue
                    previous = groups[-1][-1]
                    gap = float(span["bbox"][0]) - float(previous["bbox"][2])
                    threshold = max(
                        14.0,
                        float(span.get("size", 9.0)) * 1.5,
                        float(previous.get("size", 9.0)) * 1.5,
                    )
                    if gap > threshold:
                        groups.append([span])
                    else:
                        groups[-1].append(span)
                for group in groups:
                    shape_counter += 1
                    _append_positioned_text_group(
                        anchor,
                        shape_id=f"matchmind_text_{shape_counter}",
                        spans=group,
                        page_width=float(page.rect.width),
                    )

        if page_number < source.page_count - 1:
            document.add_page_break()

    source.close()
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _convert_pdf_to_docx_with_word(source_bytes: bytes, word_application) -> bytes | None:
    """Use an existing Word session only when fast structural conversion fails."""
    if word_application is None:
        return None
    document = None
    try:
        with tempfile.TemporaryDirectory(prefix="matchmind_pdf_word_fallback_") as directory:
            source_path = Path(directory) / "source.pdf"
            output_path = Path(directory) / "source.docx"
            source_path.write_bytes(source_bytes)
            document = word_application.Documents.Open(
                str(source_path.resolve()),
                ConfirmConversions=False,
                ReadOnly=False,
                AddToRecentFiles=False,
            )
            document.SaveAs2(str(output_path.resolve()), FileFormat=16)
            document.Close(False)
            document = None
            return output_path.read_bytes() if output_path.exists() else None
    except Exception:
        return None
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass


def _inspect_searchable_pdf(pdf_bytes: bytes, expected_pages: int | None = None) -> dict:
    """Render every page and verify searchable text, blank pages and page count."""
    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    issues: list[str] = []
    searchable_characters = 0
    for page_number, page in enumerate(document, start=1):
        text = re.sub(r"\s+", "", page.get_text("text"))
        searchable_characters += len(text)
        if not text:
            issues.append(f"Page {page_number} has no selectable text.")
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        if pixmap.width < 100 or pixmap.height < 100:
            issues.append(f"Page {page_number} could not be rendered at review size.")
    page_count = document.page_count
    document.close()
    if expected_pages is not None and page_count != expected_pages:
        issues.append(f"Expected {expected_pages} page(s), but the tailored PDF contains {page_count}.")
    if searchable_characters < 30:
        issues.append("The tailored PDF does not contain enough selectable text.")
    return {
        "passed": not issues,
        "issues": issues,
        "page_count": page_count,
        "summary": (
            f"Rendered and checked all {page_count} page(s); selectable text and page structure verified."
            if not issues
            else "Automated page review found: " + " ".join(issues)
        ),
    }


def _start_microsoft_word():
    """Start one hidden Word process that can be reused for a résumé batch."""
    if platform.system() != "Windows":
        return None
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return None

    pythoncom.CoInitialize()
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        word.ScreenUpdating = False
        return pythoncom, word
    except Exception:
        pythoncom.CoUninitialize()
        return None


def _stop_microsoft_word(session) -> None:
    if session is None:
        return
    pythoncom, word = session
    try:
        word.Quit()
    except Exception:
        pass
    pythoncom.CoUninitialize()


def _create_pdf_with_microsoft_word(
    candidate_name: str,
    docx_bytes: bytes,
    word_application=None,
) -> tuple[str, bytes] | None:
    """Use desktop Word on Windows so PDF and DOCX have the same visual layout."""
    if platform.system() != "Windows":
        return None
    owned_session = None
    word = word_application
    document = None
    try:
        if word is None:
            owned_session = _start_microsoft_word()
            if owned_session is None:
                return None
            word = owned_session[1]
        with tempfile.TemporaryDirectory(prefix="matchmind_resume_") as directory:
            docx_path = Path(directory) / "enhanced_resume.docx"
            pdf_path = Path(directory) / "enhanced_resume.pdf"
            docx_path.write_bytes(docx_bytes)
            document = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
            document.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
            return f"{_safe_filename(candidate_name)}_enhanced.pdf", pdf_path.read_bytes()
    except Exception:
        return None
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if owned_session is not None:
            _stop_microsoft_word(owned_session)


def _create_pdf_with_libreoffice(candidate_name: str, docx_bytes: bytes) -> tuple[str, bytes] | None:
    """Use LibreOffice when available so PDF and DOCX share one rendered layout."""
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return None
    try:
        with tempfile.TemporaryDirectory(prefix="matchmind_resume_") as directory:
            working_directory = Path(directory)
            docx_path = working_directory / "enhanced_resume.docx"
            pdf_path = working_directory / "enhanced_resume.pdf"
            docx_path.write_bytes(docx_bytes)
            subprocess.run(
                [
                    executable,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(working_directory),
                    str(docx_path),
                ],
                check=True,
                capture_output=True,
                timeout=90,
            )
            if not pdf_path.exists():
                return None
            return f"{_safe_filename(candidate_name)}_enhanced.pdf", pdf_path.read_bytes()
    except Exception:
        return None


def create_enhanced_resume_files(
    candidate_name: str,
    target_role: str,
    job_description: str,
    source_resume: str,
    source_name: str | None = None,
    source_bytes: bytes | None = None,
    word_application=None,
    start_word_if_needed: bool = True,
    tailoring_result: dict | None = None,
    additional_information: str = "",
    reported_gaps: str = "",
) -> dict:
    """Return truthful job-focused DOCX/PDF files without slow Word PDF import."""
    started = time.perf_counter()
    tailoring_result = tailoring_result or tailor_resume_with_ai(
        job_description,
        source_resume,
        additional_information=additional_information,
        reported_gaps=reported_gaps,
        target_role=target_role,
    )
    content = build_enhanced_resume(
        candidate_name,
        target_role,
        job_description,
        source_resume,
        tailoring_result=tailoring_result,
    )
    source_extension = Path(source_name).suffix.casefold() if source_name else ""
    preserved_docx = None
    preserved_pdf = None
    pdf_rewrite_warnings: list[str] = []
    preservation_warnings: list[str] = []
    expected_pdf_pages = None
    layout_note = "The source file did not contain a reusable visual layout."
    if source_extension in {".docx", ".pdf"} and not source_bytes:
        raise ValueError(
            "The original résumé file is unavailable. Upload it again so MatchMind can preserve its formatting."
        )

    if source_name and source_bytes and source_extension == ".docx":
        preserved_docx = _enhance_source_docx(
            candidate_name,
            source_bytes,
            content["focus_terms"],
            tailoring_result=tailoring_result,
        )
        if not preserved_docx:
            raise ValueError(
                "The uploaded DOCX could not be enhanced without losing its formatting. "
                "No generic résumé was created."
            )
        layout_note = (
            "The uploaded DOCX structure, fonts, colours, paragraph styles, tables, spacing and page settings were retained."
        )
    elif source_name and source_bytes and source_extension == ".pdf":
        text_based, expected_pdf_pages = _pdf_text_page_count(source_bytes)
        if not text_based:
            raise ValueError(
                "This PDF is image-based. Upload the editable DOCX or a text-based PDF so the résumé can remain "
                "selectable and preserve its design. No notepad-style fallback was created."
            )
        try:
            pdf_name, pdf_bytes, pdf_rewrite_warnings = _enhance_text_pdf_in_place(
                candidate_name,
                source_bytes,
                tailoring_result,
            )
            preserved_pdf = (pdf_name, pdf_bytes)
        except Exception:
            preserved_pdf = (f"{_safe_filename(candidate_name)}_enhanced.pdf", source_bytes)
            preservation_warnings.append(
                "One safe in-place PDF rewrite could not be completed, so the original designed PDF was retained "
                "instead of creating a generic résumé."
            )

        positioned_docx = _convert_pdf_to_positioned_docx(preserved_pdf[1])
        preserved_docx = (
            f"{_safe_filename(candidate_name)}_enhanced.docx",
            positioned_docx,
        )
        layout_note = (
            "The tailored PDF retains the original fonts, colours, graphics, page dimensions and page count. "
            "The editable DOCX recreates the same design with positioned, selectable text rather than page screenshots."
        )

    docx_name, docx_bytes = preserved_docx or create_enhanced_docx(candidate_name, content)
    word_pdf = None
    if word_application is not None or start_word_if_needed:
        word_pdf = _create_pdf_with_microsoft_word(
            candidate_name,
            docx_bytes,
            word_application=word_application,
        )
    office_pdf = word_pdf or _create_pdf_with_libreoffice(candidate_name, docx_bytes)
    docx_render_quality = None
    if office_pdf:
        docx_render_quality = _inspect_searchable_pdf(
            office_pdf[1],
            expected_pages=expected_pdf_pages if preserved_docx and source_extension == ".pdf" else None,
        )

    if source_extension == ".pdf" and docx_render_quality and not docx_render_quality["passed"]:
        preservation_warnings.append(
            "The editable DOCX did not exactly match the PDF page structure during automated review; its converted "
            "design was retained rather than replaced with a generic layout."
        )

    if source_extension == ".pdf" and preserved_pdf:
        pdf_name, pdf_bytes = preserved_pdf
        if tailoring_result.get("section_additions"):
            pdf_rewrite_warnings.append(
                "A verified free-text addition was not appended as a new PDF block because preserving the original layout and page count takes priority."
            )
    elif source_extension == ".docx" and office_pdf:
        pdf_name, pdf_bytes = office_pdf
    elif source_extension == ".docx":
        raise ValueError(
            "The enhanced DOCX was created with its original formatting, but a matching PDF could not be exported. "
            "Install Microsoft Word locally or include LibreOffice in the deployment. No generic PDF was created."
        )
    else:
        pdf_name, pdf_bytes = office_pdf or create_enhanced_pdf(candidate_name, content)
    pdf_bytes = _remove_blank_pdf_pages(pdf_bytes)
    quality = _inspect_searchable_pdf(
        pdf_bytes,
        expected_pages=expected_pdf_pages if source_extension == ".pdf" else None,
    )
    warning_messages = [
        message
        for message in (
            tailoring_result.get("warning"),
            *preservation_warnings,
            *pdf_rewrite_warnings,
        )
        if message
    ]
    return {
        "docx_name": docx_name,
        "docx_bytes": docx_bytes,
        "pdf_name": pdf_name,
        "pdf_bytes": pdf_bytes,
        "focused_terms": content["focus_terms"],
        "matching_layout": preserved_pdf is None and office_pdf is not None,
        "source_layout_preserved": preserved_pdf is not None or preserved_docx is not None,
        "layout_note": layout_note,
        "change_summary": tailoring_result.get("change_summary", []),
        "strong_matches": tailoring_result.get("strong_matches", []),
        "remaining_gaps": tailoring_result.get("remaining_gaps", []),
        "questions": tailoring_result.get("questions", []),
        "generation_mode": tailoring_result.get("generation_mode", "local evidence-based"),
        "warning": " ".join(warning_messages) or None,
        "qa_passed": quality["passed"],
        "qa_summary": quality["summary"],
        "qa_issues": quality["issues"],
        "elapsed_seconds": round(time.perf_counter() - started, 2),
    }


def create_enhanced_resumes_batch(
    requests: list[dict],
    progress_callback: Callable[[int, int, str, str], None] | None = None,
    maximum_workers: int = 3,
) -> tuple[dict, list[str]]:
    """Tailor resumes concurrently, then reuse one Word process for fast PDF export."""
    generated: dict = {}
    failures: list[str] = []
    total = len(requests)
    tailoring_results: dict[str, dict] = {}

    def create_tailoring_result(request: dict) -> tuple[str, dict]:
        candidate = request["candidate_name"]
        return candidate, tailor_resume_with_ai(
            request["job_description"],
            request["source_resume"],
            additional_information=request.get("additional_information", ""),
            reported_gaps=request.get("reported_gaps", ""),
            target_role=request.get("target_role", ""),
        )

    worker_count = max(1, min(maximum_workers, total or 1))
    if progress_callback:
        progress_callback(0, total, "All suitable candidates", "tailoring")
    with ThreadPoolExecutor(max_workers=worker_count) as executor:
        futures = [executor.submit(create_tailoring_result, request) for request in requests]
        for completed, future in enumerate(as_completed(futures), start=1):
            try:
                candidate, result = future.result()
                tailoring_results[candidate] = result
            except Exception as error:
                failures.append(f"Content enhancement: {error}")
            if progress_callback:
                progress_callback(completed, total, "Content enhancement", "tailoring")

    word_session = _start_microsoft_word()
    word_application = word_session[1] if word_session is not None else None

    try:
        for index, request in enumerate(requests, start=1):
            candidate = request["candidate_name"]
            if progress_callback:
                progress_callback(index - 1, total, candidate, "formatting")
            try:
                generated[candidate] = create_enhanced_resume_files(
                    candidate_name=candidate,
                    target_role=request["target_role"],
                    job_description=request["job_description"],
                    source_resume=request["source_resume"],
                    source_name=request.get("source_name"),
                    source_bytes=request.get("source_bytes"),
                    word_application=word_application,
                    start_word_if_needed=False,
                    tailoring_result=tailoring_results.get(candidate),
                    additional_information=request.get("additional_information", ""),
                    reported_gaps=request.get("reported_gaps", ""),
                )
            except Exception as error:
                failures.append(f"{candidate}: {error}")
            if progress_callback:
                progress_callback(index, total, candidate, "complete")
    finally:
        _stop_microsoft_word(word_session)

    return generated, failures

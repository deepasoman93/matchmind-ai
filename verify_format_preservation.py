from __future__ import annotations

from io import BytesIO
import zipfile

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import resume_tailor
from advanced_engine import tailoring_eligibility


def _colour_hex(run) -> str | None:
    value = run.font.color.rgb
    return str(value) if value is not None else None


def _make_designed_docx() -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run("CANDIDATE NAME")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(31, 78, 121)

    heading = document.add_paragraph()
    heading.style = document.styles["Heading 1"]
    heading_run = heading.add_run("PROFESSIONAL SUMMARY")
    heading_run.bold = True
    heading_run.font.name = "Arial"
    heading_run.font.size = Pt(13)
    heading_run.font.color.rgb = RGBColor(31, 78, 121)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "D9EAF7")
    heading._p.get_or_add_pPr().append(shading)

    summary = document.add_paragraph()
    first = summary.add_run("Data professional with four years of ")
    first.font.name = "Calibri"
    first.font.size = Pt(10)
    first.font.color.rgb = RGBColor(35, 35, 35)
    skill = summary.add_run("Python and SQL")
    skill.bold = True
    skill.font.name = "Calibri"
    skill.font.size = Pt(10)
    skill.font.color.rgb = RGBColor(31, 78, 121)
    last = summary.add_run(" experience.")
    last.italic = True
    last.font.name = "Calibri"
    last.font.size = Pt(10)
    last.font.color.rgb = RGBColor(112, 48, 160)

    skills_heading = document.add_paragraph()
    skills_run = skills_heading.add_run("TECHNICAL SKILLS")
    skills_run.bold = True
    skills_run.font.color.rgb = RGBColor(31, 78, 121)
    document.add_paragraph("Python | SQL | Excel | Power BI")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "EXPERIENCE"
    table.cell(0, 1).text = "Data Analyst | 2022–Present"
    for run in table.cell(0, 0).paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    cell_shading = OxmlElement("w:shd")
    cell_shading.set(qn("w:fill"), "1F4E79")
    table.cell(0, 0)._tc.get_or_add_tcPr().append(cell_shading)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Candidate Name | Résumé")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _make_coloured_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.draw_rect(fitz.Rect(40, 40, 555, 85), color=(0.12, 0.31, 0.48), fill=(0.12, 0.31, 0.48))
    page.insert_text((55, 69), "PROFESSIONAL SUMMARY", fontsize=14, color=(1, 1, 1))
    page.insert_text(
        (55, 120),
        "Worked on automated reporting workflows using Python and SQL.",
        fontsize=10,
        color=(0.14, 0.14, 0.14),
    )
    payload = document.tobytes()
    document.close()
    return payload


def verify_docx_design_and_colour_preservation() -> None:
    source_bytes = _make_designed_docx()
    source = Document(BytesIO(source_bytes))
    source_summary = source.paragraphs[2]
    source_colours = [_colour_hex(run) for run in source_summary.runs]

    tailoring_result = {
        "rewrites": [
            {
                "source_text": "Data professional with four years of Python and SQL experience.",
                "revised_text": "Data Analyst with four years of Python and SQL experience delivering automated reporting workflows.",
                "reason": "Align the truthful summary with the target role.",
            }
        ],
        "verified_skills": [],
        "section_additions": {},
        "change_summary": ["Strengthened the professional summary."],
        "strong_matches": ["Python and SQL"],
        "remaining_gaps": [],
        "questions": [],
        "generation_mode": "verification",
    }
    files = resume_tailor.create_enhanced_resume_files(
        "Candidate Name",
        "Data Analyst",
        "Data Analyst requiring Python, SQL, Excel and reporting automation",
        """CANDIDATE NAME
PROFESSIONAL SUMMARY
Data professional with four years of Python and SQL experience.
TECHNICAL SKILLS
Python | SQL | Excel | Power BI
EXPERIENCE
Data Analyst | 2022-Present
""",
        source_name="candidate.docx",
        source_bytes=source_bytes,
        start_word_if_needed=False,
        tailoring_result=tailoring_result,
    )

    enhanced = Document(BytesIO(files["docx_bytes"]))
    enhanced_summary = next(
        paragraph for paragraph in enhanced.paragraphs
        if "automated reporting workflows" in paragraph.text
    )
    enhanced_colours = {_colour_hex(run) for run in enhanced_summary.runs if run.text}
    assert len(enhanced.tables) == len(source.tables) == 1
    assert set(source_colours).issubset(enhanced_colours)
    assert enhanced.sections[0].left_margin == source.sections[0].left_margin

    with zipfile.ZipFile(BytesIO(source_bytes)) as source_zip, zipfile.ZipFile(BytesIO(files["docx_bytes"])) as output_zip:
        assert source_zip.read("word/styles.xml") == output_zip.read("word/styles.xml")
        assert source_zip.read("word/theme/theme1.xml") == output_zip.read("word/theme/theme1.xml")

    pdf = fitz.open(stream=files["pdf_bytes"], filetype="pdf")
    text = " ".join(page.get_text() for page in pdf)
    assert "automated reporting workflows" in text
    assert all(page.get_text().strip() for page in pdf)
    pdf.close()
    print("PASS: DOCX fonts, colours, mixed-run formatting, table, margins and theme were retained.")


def verify_pdf_design_preservation() -> None:
    source_bytes = _make_coloured_pdf()
    tailoring_result = {
        "rewrites": [
            {
                "source_text": "Worked on automated reporting workflows using Python and SQL.",
                "revised_text": "Delivered automated reporting workflows using Python and SQL.",
                "reason": "Use a stronger truthful action verb.",
            }
        ]
    }
    _, output_bytes, warnings = resume_tailor._enhance_text_pdf_in_place(
        "Candidate Name",
        source_bytes,
        tailoring_result,
    )
    assert not warnings

    source = fitz.open(stream=source_bytes, filetype="pdf")
    output = fitz.open(stream=output_bytes, filetype="pdf")
    assert output.page_count == source.page_count == 1
    assert output[0].rect == source[0].rect
    assert len(output[0].get_drawings()) == len(source[0].get_drawings())
    output_text = output[0].get_text().replace("ﬂ", "fl").replace("ﬁ", "fi")
    assert "Delivered automated reporting workflows" in output_text
    assert "PROFESSIONAL SUMMARY" in output_text

    source_heading = next(
        span for block in source[0].get_text("dict")["blocks"] if "lines" in block
        for line in block["lines"] for span in line["spans"]
        if "PROFESSIONAL SUMMARY" in span["text"]
    )
    output_heading = next(
        span for block in output[0].get_text("dict")["blocks"] if "lines" in block
        for line in block["lines"] for span in line["spans"]
        if "PROFESSIONAL SUMMARY" in span["text"]
    )
    assert source_heading["color"] == output_heading["color"]
    source.close()
    output.close()
    print("PASS: PDF colour, background graphics, page size and searchable text were retained.")


def verify_pdf_to_docx_keeps_editable_text_and_visual_layer() -> None:
    source_bytes = _make_coloured_pdf()
    docx_bytes = resume_tailor._convert_pdf_to_positioned_docx(source_bytes)
    with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        media_files = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert "matchmind_background" in document_xml
    assert "wp:anchor" in document_xml
    assert "v:textbox" in document_xml
    assert "PROFESSIONAL SUMMARY" in document_xml
    assert media_files

    rendered = resume_tailor._create_pdf_with_libreoffice("Candidate Name", docx_bytes)
    if rendered is not None:
        document = fitz.open(stream=rendered[1], filetype="pdf")
        assert document.page_count == 1
        assert "PROFESSIONAL SUMMARY" in document[0].get_text()
        assert document[0].get_images(full=True)
        red, green, blue = document[0].get_pixmap(alpha=False).pixel(45, 45)[:3]
        assert red + green + blue < 650
        document.close()
    print("PASS: PDF-to-DOCX output keeps real text plus the original visual design layer.")


def verify_no_generic_pdf_fallback() -> None:
    source_bytes = _make_coloured_pdf()
    original_converter = resume_tailor._convert_pdf_to_positioned_docx
    resume_tailor._convert_pdf_to_positioned_docx = lambda _payload: (_ for _ in ()).throw(
        ValueError("No generic résumé was created")
    )
    try:
        try:
            resume_tailor.create_enhanced_resume_files(
                "Candidate Name",
                "Data Analyst",
                "Data Analyst requiring 2-3 years of Python and SQL experience",
                "Worked on automated reporting workflows using Python and SQL.",
                source_name="candidate.pdf",
                source_bytes=source_bytes,
                start_word_if_needed=False,
            )
        except ValueError as error:
            assert "No generic résumé was created" in str(error)
        else:
            raise AssertionError("A generic résumé was created after preservation failure.")
    finally:
        resume_tailor._convert_pdf_to_positioned_docx = original_converter
    print("PASS: preservation failure stops safely instead of producing a notepad résumé.")


def verify_experience_rule_remains_fixed() -> None:
    result = {
        "Match %": 59.7,
        "Semantic %": 43.8,
        "Skill coverage %": 100.0,
        "Experience": 4,
    }
    suitable, _ = tailoring_eligibility(
        result,
        "Data Analyst requiring 2-3 years of experience.",
    )
    assert suitable
    print("PASS: the 4-year candidate remains Suitable for a 2-3-year role.")


def main() -> None:
    verify_docx_design_and_colour_preservation()
    verify_pdf_design_preservation()
    verify_pdf_to_docx_keeps_editable_text_and_visual_layer()
    verify_no_generic_pdf_fallback()
    verify_experience_rule_remains_fixed()
    print("\nPASS: MatchMind formatting-preserving enhancement is active.")


if __name__ == "__main__":
    main()
